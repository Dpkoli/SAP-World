from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "public-launch"
OUT.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5F6B7A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GREY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
RED = "9B1C1C"
GREEN = "276749"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int], indent: int = 120) -> None:
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"Table widths must total {TABLE_WIDTH}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_hyperlink(paragraph, text: str, url: str, color=BLUE) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    r_pr.append(r_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_font(run, size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_document(doc: Document, running_left: str, running_right: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    for sec in doc.sections:
        sec.different_first_page_header_footer = True
        header = sec.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(running_left)
        set_font(r, 9, MUTED, bold=True)
        r2 = p.add_run("  |  " + running_right)
        set_font(r2, 9, MUTED)
        footer = sec.footer
        add_page_field(footer.paragraphs[0])


def cover(doc: Document, kicker: str, title: str, subtitle: str, audience: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(124)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(kicker.upper())
    set_font(r, 10.5, GOLD, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_font(r, 29, NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(30)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    set_font(r, 15, DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(78)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A practical, step-by-step field guide")
    set_font(r, 10.5, GOLD, italic=True)

    for text, size, bold in (
        ("Version 1.0 | 30 June 2026", 11, True),
        (audience, 9.5, False),
        ("Live app: https://sap-world-pi.vercel.app", 9.5, False),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, size, NAVY if bold else MUTED, bold=bold)
    doc.add_page_break()


def para(doc: Document, text: str, bold_prefix: str | None = None, style=None) -> None:
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbers(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def callout(doc: Document, label: str, text: str, fill=CALLOUT, label_color=DARK_BLUE) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "6")
        border.set(qn("w:color"), "D7DBE2")
        borders.append(border)
    p_pr.append(borders)
    r = p.add_run(label + ": ")
    set_font(r, 10.5, label_color, bold=True)
    r = p.add_run(text)
    set_font(r, 10.5, NAVY)


def app_ref(doc: Document, path: str) -> None:
    callout(doc, "App reference", path, LIGHT_BLUE, BLUE)


def try_it(doc: Document, text: str) -> None:
    callout(doc, "Try it now", text, "EEF7F1", GREEN)


def warning(doc: Document, text: str) -> None:
    callout(doc, "Important", text, "FFF5F5", RED)


def table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    set_table_geometry(t, widths)
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        shade_cell(hdr.cells[idx], LIGHT_BLUE)
        p = hdr.cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(header) < 16 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        set_font(r, 9.5, NAVY, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 and len(value) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_font(r, 9.5, NAVY)
    set_table_geometry(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def chapter(doc: Document, title: str, lead: str | None = None) -> None:
    doc.add_heading(title, level=1)
    if lead:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(lead)
        set_font(r, 11.5, DARK_BLUE, italic=True)


def new_chapter(doc: Document, title: str, lead: str | None = None) -> None:
    doc.add_page_break()
    chapter(doc, title, lead)


def build_user_manual() -> Path:
    doc = Document()
    configure_document(doc, "SAP World User Manual", "Learner and Administrator Guide")
    cover(
        doc,
        "SAP World",
        "User Manual & Learning Guide",
        "Learn the app while learning how connected SAP work flows through an enterprise",
        "For learners, reviewers, trainers, and administrators",
    )

    chapter(doc, "How to use this guide", "Use the guide beside the app. Read one short section, perform the action, then check the result on screen.")
    para(doc, "Each practical explanation contains an App reference. It names the exact menu or control to open. A Try it now box gives you a small action to complete before moving on.")
    callout(doc, "Recommended rhythm", "Learn one process at a time: understand the business purpose, complete the guided transaction, diagnose an exception, record evidence, and review readiness.", LIGHT_BLUE)
    doc.add_heading("Who the app is for", level=2)
    bullets(doc, [
        "Learners who want practical SAP S/4HANA process understanding without needing access to a live SAP tenant.",
        "Career changers who need a structured route through SAP modules, transactions, evidence, and business language.",
        "Managers and reviewers who want to inspect learning evidence, simulated approvals, and readiness.",
        "Administrators who manage identities, controlled content, certification reviews, and release readiness.",
    ])
    doc.add_heading("What SAP World is - and is not", level=2)
    para(doc, "SAP World is a learning simulation. It uses realistic enterprise data, document chains, approvals, ledgers, controls, and guided explanations. It is not a live SAP production system, official SAP training, or a guarantee of employment or certification.")
    warning(doc, "Do not enter real employer data, customer details, passwords, confidential SAP screenshots, or personal information into practice notes or SAP Mentor prompts.")

    doc.add_heading("Contents", level=2)
    table(doc, ["Part", "What you will learn"], [
        ("1", "Sign in safely and understand the workspace"),
        ("2", "Follow a complete learner journey"),
        ("3", "Use every main learner area"),
        ("4", "Build evidence, readiness, and capstone submissions"),
        ("5", "Run simulations and interpret enterprise results"),
        ("6", "Use SAP Mentor responsibly"),
        ("7", "Use administrator and reviewer controls"),
        ("8", "Troubleshoot common issues and follow a study plan"),
    ], [1200, 8160])

    new_chapter(doc, "1. Start safely: account, sign-in, and recovery", "Your account keeps progress, evidence, simulation history, and role-based access connected to you.")
    doc.add_heading("Create a learner account", level=2)
    numbers(doc, [
        "Open https://sap-world-pi.vercel.app.",
        "Choose the registration option on the sign-in screen.",
        "Enter the requested account details and use a unique password that you do not use elsewhere.",
        "Submit the form, then sign in with the new account.",
        "Confirm that the Enterprise overview opens and your learner navigation appears.",
    ])
    app_ref(doc, "Opening screen -> Register / Create account -> Sign in -> Enterprise overview")
    try_it(doc, "Sign out once and sign in again. This confirms that you know your credentials before starting a long learning session.")

    doc.add_heading("Recover a password", level=2)
    numbers(doc, [
        "Choose Forgot password on the sign-in screen.",
        "Enter the email address used for the account.",
        "Follow the reset instructions. A reset token is time-limited and can be used only once.",
        "Create a new unique password and sign in again.",
    ])
    app_ref(doc, "Opening screen -> Forgot password")
    warning(doc, "A password reset ends existing sessions. If you did not request the reset, contact the site operator and change your password immediately.")

    doc.add_heading("Understand roles", level=2)
    table(doc, ["Role", "What appears in the app", "Use it for"], [
        ("Learner", "Learning, simulations, tutor, evidence, analytics, history, and approvals", "Practising connected SAP work"),
        ("Administrator", "All learner areas plus Control plane", "Identity, governance, reviews, readiness, and release operations"),
        ("Suspended account", "Access is blocked", "Safety or lifecycle control; an administrator must reactivate it"),
    ], [1500, 3900, 3960])

    new_chapter(doc, "2. Your first 30 minutes", "This short session teaches the navigation and gives you a first piece of saved progress.")
    table(doc, ["Minutes", "Action", "App reference"], [
        ("0-5", "Review the dashboard and choose a focus process.", "Enterprise overview"),
        ("5-10", "Read one role-based pathway and its business outcome.", "Learning centre"),
        ("10-20", "Open the guided tutor and complete the first transaction step.", "Transaction tutor -> Guided"),
        ("20-25", "Write a short evidence note: document, key field, effect, and proof.", "Transaction tutor -> Evidence note"),
        ("25-30", "Ask one grounded integration question and review your next action.", "Ask SAP Mentor / Learning centre"),
    ], [1000, 4700, 3660])
    try_it(doc, "Use a simple evidence note such as: 'Posted goods receipt for the purchase order; verified material, quantity, plant, and movement; stock increased and the GR/IR accounting effect was created.'")

    doc.add_heading("The learning loop", level=2)
    numbers(doc, [
        "Orient: understand the enterprise context and the business reason for the process.",
        "Observe: read the transaction playbook and document chain.",
        "Perform: complete each guided transaction step.",
        "Explain: record why the step matters and what changed downstream.",
        "Recover: solve the troubleshooting diagnostic.",
        "Prove: submit evidence and, when open, complete the capstone.",
        "Reflect: review readiness and choose the next weakest area.",
    ])
    app_ref(doc, "Enterprise overview -> Learning centre -> Transaction tutor -> Troubleshoot / Capstone -> Portfolio")

    new_chapter(doc, "3. Navigate the learner workspace", "The left navigation follows the way an SAP learner moves from context to practice and proof.")
    nav_rows = [
        ("Enterprise overview", "See the business snapshot, portfolio progress, focus process, and urgent actions.", "Start every session here."),
        ("Learning centre", "Choose role-based pathways, review progress, readiness, portfolio, and certification evidence.", "Use this as your learning home."),
        ("Industry blueprints", "Compare ten industries and see value chains, master data, KPIs, compliance, and failures.", "Use before choosing a specialism."),
        ("Simulation studio", "Generate a deterministic multi-year SAP scenario from industry, year, event, and volume choices.", "Use to create connected practice data."),
        ("Process explorer", "Trace document flow, upstream/downstream links, inventory, workflow, and journal evidence.", "Use to understand integration."),
        ("Advanced transactions", "Practise specialist labs such as stock transfer, returns, assets, VAT, and year-end close.", "Use after core pathways."),
        ("Performance analytics", "Interpret revenue, margin, inventory, service, downtime, waste, and profitability movements.", "Use to connect SAP work to management outcomes."),
        ("Simulation history", "Review the enterprise chronology and how documents and events changed results.", "Use for revision and audit thinking."),
        ("Approval inbox", "Inspect evidence and make simulated approve, reject, or request-information decisions.", "Use to practise controls."),
        ("Transaction tutor", "Work step by step through guided, troubleshooting, implementation, and capstone modes.", "Use for hands-on learning."),
    ]
    table(doc, ["Area", "What it teaches", "Best use"], nav_rows, [2200, 4550, 2610])
    app_ref(doc, "Left navigation -> select any learner area listed above")

    doc.add_heading("Enterprise data areas", level=2)
    table(doc, ["Area", "Question it answers"], [
        ("Company structure", "How do legal entities, sales, purchasing, manufacturing, and inventory units connect?"),
        ("Master data", "Which materials, valuation, MRP, sourcing, batch, quality, BOM, routing, and work-centre records drive transactions?"),
        ("Data governance", "How are master-data changes validated, approved, versioned, and audited?"),
        ("Plants & operations", "Where are capacity, utilisation, staffing, storage, and operational demand visible?"),
        ("Business partners", "Which suppliers and customers carry commercial exposure, risk, or blocks?"),
    ], [2400, 6960])
    try_it(doc, "Open Company structure, select a plant, then open Master data and find a material associated with that operational context.")

    new_chapter(doc, "4. Learn the eight core SAP processes", "Each pathway teaches a business outcome, SAP actions, document evidence, integration effects, and exception recovery.")
    processes = [
        ("P2P", "Procure-to-Pay", "Purchase requisition to supplier payment", "MM, QM, FI"),
        ("O2C", "Order-to-Cash", "Customer demand to delivery and cash", "SD, EWM, PP, FI"),
        ("P2P/PP", "Plan-to-Produce", "MRP demand to production completion", "PP, MM, EWM, CO, FI"),
        ("R2R", "Record-to-Report", "Operational postings to period close and reporting", "FI-GL, AP, AR, CO"),
        ("QM", "Quality Management", "Inspection lot to stock disposition and supplier feedback", "QM, MM"),
        ("PM", "Maintenance", "Breakdown to repair, cost, safety, and reliability learning", "PM, MM, CO"),
        ("H2R", "Hire-to-Retire", "Onboarding through payroll and organisational assignment", "HCM, SuccessFactors, FI, CO"),
        ("W2D", "Warehouse-to-Dispatch", "Warehouse demand to picking, goods issue, and transport", "EWM, TM, SD, MM, FI"),
    ]
    table(doc, ["Code", "Pathway", "Business chain", "Modules"], processes, [900, 2200, 4000, 2260])
    app_ref(doc, "Learning centre -> choose a pathway -> Open lesson / Continue learning")

    doc.add_heading("How to complete a guided lesson", level=2)
    numbers(doc, [
        "Choose one pathway in Learning centre and read the outcome, role, modules, and prerequisites.",
        "Open Transaction tutor and select Guided mode.",
        "For the active step, read the prerequisite, SAP app or transaction, fields, validation, integration impact, and proof to keep.",
        "Perform the simulated action and inspect the resulting document or status.",
        "Write an evidence note with four parts: document proof, key-field proof, integration effect, and audit detail.",
        "Use the evidence coach feedback to improve the note before continuing.",
        "Complete the knowledge check and review corrective feedback.",
    ])
    app_ref(doc, "Transaction tutor -> Guided -> Active step -> SAP processing guide -> Evidence note -> Knowledge check")
    try_it(doc, "Before clicking the next step, explain aloud: What was the input? What SAP object changed? Which process receives the output next?")

    doc.add_heading("How to read a document chain", level=2)
    para(doc, "A document chain is the evidence that one business action caused the next. In P2P, for example, the chain moves from purchase requisition to purchase order, goods receipt, quality inspection, invoice verification, and vendor payment.")
    app_ref(doc, "Process explorer -> select a process/document -> Upstream links / Downstream links / Inventory / Journal evidence")
    bullets(doc, [
        "Upstream: what authorised or created the current document?",
        "Current document: what business event was recorded?",
        "Downstream: which operational or financial action is now possible?",
        "Control: which approval, match, tolerance, or validation protects the step?",
        "Evidence: which document number, status, quantity, amount, or journal proves completion?",
    ])

    new_chapter(doc, "5. Troubleshooting, implementation, and capstone modes", "Knowing the happy path is useful; knowing how to diagnose failure is what makes the learning job-ready.")
    doc.add_heading("Troubleshooting mode", level=2)
    numbers(doc, [
        "Open the same pathway in Transaction tutor and switch to Troubleshoot.",
        "Read the symptom without jumping directly to the answer.",
        "Identify the failed document, status, master data, control, or integration hand-off.",
        "Choose the diagnosis and recovery action.",
        "Review corrective feedback and retry if required.",
        "Confirm that diagnostic completion appears in your progress and readiness.",
    ])
    app_ref(doc, "Transaction tutor -> Troubleshoot -> Select diagnosis / recovery -> Submit")
    try_it(doc, "Use the sequence 'symptom -> evidence -> likely cause -> safe correction -> downstream check'. Write one sentence for each before submitting.")

    doc.add_heading("Implementation mode", level=2)
    para(doc, "Implementation mode changes the question from 'how do I post this transaction?' to 'what must be designed, configured, integrated, tested, controlled, and supported so the process works reliably?'")
    app_ref(doc, "Transaction tutor -> Implementation")
    bullets(doc, [
        "Organisation: company code, plant, purchasing, sales, storage, and responsibility design.",
        "Master data: ownership, validation, dependencies, migration, and cutover.",
        "Configuration: process rules, tolerances, account determination, outputs, and controls.",
        "Integration: hand-offs between modules and evidence that postings reconcile.",
        "Testing: positive, negative, volume, security, and end-to-end scenarios.",
        "Go-live: access, monitoring, support, fallback, and hypercare.",
    ])

    doc.add_heading("Capstone mode", level=2)
    para(doc, "A capstone opens only when the required learning evidence is ready. It asks you to explain the process, document chain, control logic, exception handling, and business impact in your own words.")
    app_ref(doc, "Transaction tutor -> Capstone; Learning centre -> Portfolio / Readiness")
    numbers(doc, [
        "Check the readiness actions for the selected process.",
        "Complete missing guided steps, diagnostics, or evidence notes.",
        "Open the capstone and read the scenario and rubric.",
        "Write a structured response of at least the minimum length shown by the app.",
        "Include document evidence, control reasoning, integration effects, and the recovery or recommendation.",
        "Submit, review the score and status, and revise if the evidence is not yet review-ready.",
    ])
    warning(doc, "Do not paste confidential employer work, copyrighted training answers, or real system credentials into a capstone response.")

    new_chapter(doc, "6. Portfolio, readiness, and certification evidence", "The portfolio converts activity into a practical record of what you can explain and prove.")
    doc.add_heading("Read your readiness score", level=2)
    para(doc, "Readiness is based on demonstrated progress: guided completion, troubleshooting diagnostics, evidence coverage, and capstone outcomes. It is not a formal SAP certification result.")
    app_ref(doc, "Enterprise overview -> Portfolio / Focus process; Learning centre -> Portfolio / Readiness")
    table(doc, ["Signal", "Meaning", "Next action"], [
        ("Guided progress", "You completed the core transaction flow.", "Finish missing tutor steps and the knowledge check."),
        ("Diagnostic progress", "You demonstrated exception diagnosis and recovery.", "Complete the troubleshooting lab."),
        ("Evidence coverage", "Your notes contain useful completion proof.", "Add document, key field, impact, and audit detail."),
        ("Capstone status", "Your scenario response has been scored and classified.", "Revise weak evidence or wait for assessor review."),
        ("Next best action", "The app has identified the highest-value next move.", "Follow it before starting a new pathway."),
    ], [1900, 3900, 3560])

    doc.add_heading("Generate certification-style evidence", level=2)
    numbers(doc, [
        "Open the portfolio in Learning centre.",
        "Review process-by-process evidence gaps and capstone status.",
        "Generate the certification evidence export when the app shows sufficient evidence.",
        "Check that the export summarises progress without exposing raw private notes or capstone text.",
        "If your organisation uses assessor review, wait for approve, reject, or revision-request status.",
    ])
    app_ref(doc, "Learning centre -> Portfolio -> Certification evidence / Export")
    warning(doc, "The export is evidence from SAP World, not an SAP SE credential. Present it as a learning portfolio unless an authorised issuer explicitly confirms otherwise.")

    new_chapter(doc, "7. Industry blueprints and choosing a specialism", "Use industry context to understand why the same SAP module is configured and controlled differently across businesses.")
    para(doc, "The catalogue contains ten industry blueprints. Each blueprint connects value chain, lifecycle, organisation, master data, KPIs, compliance, reporting, dependencies, common failures, and seasonality.")
    app_ref(doc, "Industry blueprints -> Filter by industry or SAP capability -> Open blueprint")
    numbers(doc, [
        "Open Industry blueprints and scan the industries before filtering.",
        "Choose an industry that matches your experience or target role.",
        "Read the value chain first; it explains how the business earns and delivers value.",
        "Review master data and compliance next; they reveal what must be accurate and controlled.",
        "Finish with KPIs and failures; these become useful interview and troubleshooting examples.",
        "Save the roadmap preference when the app offers it.",
    ])
    try_it(doc, "Compare one common capability, such as procurement or finance, across two industries. Record three differences in master data, controls, or KPIs.")

    new_chapter(doc, "8. Simulation Studio and generated ledgers", "Simulations let you see how connected SAP documents accumulate over time and change enterprise performance.")
    doc.add_heading("Create a simulation", level=2)
    numbers(doc, [
        "Open Simulation studio.",
        "Choose an industry, fiscal-year starting point, curated business event, and volume profile.",
        "Review the scenario summary, organisation dependencies, master data, document flow, controls, and expected impacts.",
        "Generate and save the simulation to your account.",
        "Open the resulting ledger or history and inspect chronology, process, module, exception, and journal links.",
    ])
    app_ref(doc, "Simulation studio -> Industry / Fiscal year / Event / Volume -> Generate / Save")
    para(doc, "The same inputs produce a stable scenario signature. This makes the simulation repeatable and lets the app reuse saved generated ledger evidence.")
    warning(doc, "A simulation result is educational and deterministic. Do not use its financial or operational figures as real business forecasts.")

    doc.add_heading("Interpret Performance analytics", level=2)
    app_ref(doc, "Performance analytics -> choose period, process, module, industry, or simulation")
    bullets(doc, [
        "Movement: what changed in revenue, margin, service, inventory, downtime, waste, working capital, or profitability?",
        "Driver: which event or SAP document explains the movement?",
        "Module: where was the operational or accounting effect recorded?",
        "Action: which management response is recommended, and what evidence supports it?",
        "Integrity: do document links, journals, quantities, and statuses reconcile?",
    ])

    doc.add_heading("Use Simulation history", level=2)
    para(doc, "History shows the enterprise chronology from 2023 to 2026 and links decisions, disruptions, and SAP documents to later effects. Use it to practise audit narratives: what happened, when, why, and with what result.")
    app_ref(doc, "Simulation history -> filter events/documents -> open an event -> follow related documents")

    new_chapter(doc, "9. Process explorer, approvals, and advanced labs", "These areas teach cross-module integration and control judgement.")
    doc.add_heading("Process explorer", level=2)
    numbers(doc, [
        "Choose a process and open a document in the flow.",
        "Trace the upstream source and downstream references.",
        "Check quantity, status, inventory movement, approval, and journal evidence.",
        "Explain whether the chain is complete and internally consistent.",
    ])
    app_ref(doc, "Process explorer -> Process filter -> Document -> Links / Workflow / Inventory / Journal")

    doc.add_heading("Approval inbox", level=2)
    numbers(doc, [
        "Open a pending item and read the business request.",
        "Inspect supporting evidence, amount, risk, route, and prior decisions.",
        "Choose approve only when evidence and authority are sufficient.",
        "Choose request information when a correctable evidence gap exists.",
        "Choose reject when the request is unsafe, invalid, or outside policy.",
        "Write a clear reason; the decision is retained in the audit trail.",
    ])
    app_ref(doc, "Approval inbox -> select item -> Evidence / Route -> Approve, Reject, or Request information")

    doc.add_heading("Advanced transactions", level=2)
    para(doc, "Advanced labs extend core process knowledge into stock transfers, customer returns, asset acquisition and depreciation, VAT adjustments, year-end close, and other specialist scenarios. Each lab explains the SAP entry point, fields, document effects, accounting, controls, and learner evidence.")
    app_ref(doc, "Advanced transactions -> choose lab -> complete steps -> review evidence and resulting documents")
    try_it(doc, "Complete an advanced lab only after you can explain its core pathway. For example, understand O2C before customer returns and R2R before year-end close.")

    new_chapter(doc, "10. Use SAP Mentor well", "SAP Mentor is most useful when you ask focused questions grounded in the current simulation or your saved learning progress.")
    app_ref(doc, "Bottom/right floating button -> Ask SAP Mentor")
    doc.add_heading("Good questions", level=2)
    bullets(doc, [
        "Why does this goods receipt affect inventory and GR/IR?",
        "What downstream documents should exist after this sales order step?",
        "Which evidence is missing from my P2P portfolio?",
        "Help me improve the control reasoning in my capstone evidence.",
        "What master-data issue could cause this production exception?",
    ])
    doc.add_heading("Questions to avoid", level=2)
    bullets(doc, [
        "Requests containing real customer, employee, financial, or employer-confidential data.",
        "Requests for passwords, credentials, security bypasses, or unauthorised access.",
        "Assuming an answer is official SAP, legal, accounting, tax, employment, or audit advice.",
        "Copying an answer into evidence without checking it against the simulation and document chain.",
    ])
    try_it(doc, "Use the pattern: 'In [process], at [step], I see [document/status]. Explain [integration/control question] and tell me what evidence to verify.'")

    new_chapter(doc, "11. Administrator and assessor guide", "Administration is role-protected. It controls access, content, evidence review, readiness, and production operations.")
    warning(doc, "Only authorised administrators should use the Control plane. Never promote a learner to administrator simply to resolve a learning issue.")
    app_ref(doc, "Left navigation -> Admin -> Control plane")

    doc.add_heading("Identity management", level=2)
    bullets(doc, [
        "Assign or remove administrator and learner roles using least privilege.",
        "Suspend an account when access must stop; suspension revokes active sessions.",
        "Reactivate only after verifying the person and the reason for restoration.",
        "Review access lifecycle history for assignment, removal, suspension, recovery, and reactivation events.",
        "Configure a real OIDC or SAML provider before enterprise production use.",
    ])
    app_ref(doc, "Control plane -> Identity management / Access lifecycle")

    doc.add_heading("Certification review", level=2)
    numbers(doc, [
        "Open the certification review queue.",
        "Select the learner and process submission.",
        "Review readiness, evidence coverage, diagnostic completion, capstone score, and verification notes.",
        "Record approve, reject, or revision requested with a specific assessor note.",
        "Confirm the decision appears in persistent history.",
    ])
    app_ref(doc, "Control plane -> Certification review / Assessment portfolio review")

    doc.add_heading("Content and release governance", level=2)
    bullets(doc, [
        "Use the controlled content register to inspect owners, versions, readiness, evidence gates, blockers, and release notes.",
        "Use Deployment gate summary to review storage, security, identity, content, ledger, mentor, and deployment checks.",
        "Record a release decision only for the current readiness fingerprint; changed evidence requires a new decision.",
        "Use release operations for controlled promotion or rollback and preserve the audit evidence.",
        "Treat the development roadmap as planning information, not automatic deployment approval.",
    ])
    app_ref(doc, "Control plane -> Release readiness register / Deployment gate summary / Release governance / Release operations")

    new_chapter(doc, "12. Troubleshooting and safe use", "Most learner problems can be solved by checking account state, saved progress, and the active process before repeating an action.")
    table(doc, ["Problem", "What to check", "Safe action"], [
        ("Cannot sign in", "Email, password, account suspension, reset status", "Use password recovery; contact support if suspended."),
        ("Progress did not update", "Signed-in account, network, page response", "Refresh once; reopen the lesson; do not repeatedly submit."),
        ("Capstone is locked", "Guided, diagnostic, and evidence readiness", "Complete the listed readiness actions."),
        ("Mentor answer seems wrong", "Current process, document, and cited context", "Ask a narrower question and verify against Process explorer."),
        ("Simulation not found", "Correct account and saved simulation", "Return to Simulation studio and check saved history."),
        ("Admin area missing", "Your current role", "Ask an existing administrator; do not create a second account to bypass access."),
        ("Unexpected error", "Time, page, action, and non-sensitive screenshot", "Record the details, avoid including personal data, and contact support."),
    ], [1900, 3300, 4160])
    warning(doc, "If the site shows a security warning, exposes another user's data, or behaves as though you have the wrong role, stop using it and report the issue immediately.")

    doc.add_heading("A four-week learning plan", level=2)
    table(doc, ["Week", "Goal", "App work", "Evidence of success"], [
        ("1", "Understand enterprise and one core process", "Overview, structure, master data, guided tutor", "Completed guided flow and evidence notes"),
        ("2", "Diagnose and recover", "Troubleshoot, process explorer, approvals", "Completed diagnostic and clear control reasoning"),
        ("3", "Connect operations to results", "Simulation studio, analytics, history", "Explained three KPI movements with document evidence"),
        ("4", "Prove and present", "Capstone, portfolio, certification export", "Review-ready capstone and a clear learning narrative"),
    ], [900, 2200, 3400, 2860])

    new_chapter(doc, "13. Glossary", "Use these plain-language definitions while working through the app.")
    glossary = [
        ("Document chain", "Linked SAP records showing how one business action leads to the next."),
        ("Master data", "Relatively stable business records, such as materials, suppliers, customers, BOMs, and work centres, that transactions depend on."),
        ("Transaction", "A business action recorded in an SAP app or transaction code."),
        ("Posting", "A recorded operational or financial effect, often creating a material or accounting document."),
        ("Journal evidence", "Debit and credit entries that prove the financial effect of an event."),
        ("Upstream / downstream", "The source that came before a document and the process or document that follows it."),
        ("Control", "A rule, validation, approval, match, tolerance, or segregation that reduces error or misuse."),
        ("Readiness", "SAP World's evidence-based estimate of learning progress, not an official professional certification."),
        ("Capstone", "A scenario-based assessment that asks the learner to explain and evidence an end-to-end process."),
        ("Scenario signature", "A stable identifier derived from simulation choices so results can be reproduced and reused."),
        ("OIDC / SAML", "Enterprise identity standards used to connect an organisation's sign-in provider."),
    ]
    table(doc, ["Term", "Simple meaning"], glossary, [2300, 7060])

    doc.add_heading("Finish each session with this checklist", level=2)
    bullets(doc, [
        "I can explain the business purpose of the work I completed.",
        "I can identify the document or status that proves completion.",
        "I can describe at least one upstream and one downstream effect.",
        "I recorded evidence without real confidential or personal data.",
        "I reviewed corrective feedback and my next best action.",
        "I signed out if I used a shared device.",
    ])
    app_ref(doc, "Enterprise overview -> Next best action; account menu -> Sign out")

    path = OUT / "SAP_World_User_Manual_and_Learning_Guide.docx"
    doc.save(path)
    return path


def source(doc: Document, title: str, url: str, note: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    add_hyperlink(p, title, url)
    r = p.add_run(" - " + note)
    set_font(r, 9.5, NAVY)


def build_launch_handbook() -> Path:
    doc = Document()
    configure_document(doc, "SAP World Public Launch & Monetisation", "UK-first Operating Handbook")
    cover(
        doc,
        "SAP World",
        "Public Launch & Monetisation Handbook",
        "A UK-first route from live preview to safe, public, revenue-ready service",
        "For the product owner, administrator, legal adviser, and launch partners",
    )

    chapter(doc, "Executive decision", "The site can be viewed publicly today, but it should not accept real users or payments until the P0 launch gates below are closed.")
    callout(doc, "Current state", "The Vercel production deployment is live at https://sap-world-pi.vercel.app. The latest verified health response reported local-file storage with durable=false. That storage can disappear or diverge in a serverless environment, so production PostgreSQL is the first launch blocker.", "FFF5F5", RED)
    para(doc, "This handbook is an operational plan, not legal, tax, accounting, cybersecurity, or trademark advice. Use a UK solicitor, accountant/tax adviser, and security professional before taking payments or entering new countries.")

    doc.add_heading("Go / no-go rule", level=2)
    table(doc, ["Gate", "Minimum evidence", "Decision"], [
        ("P0 - must close", "Durable database; tested backup/restore; legal pages; support contact; incident plan; security tests; payment design; SAP trademark review", "No public user acquisition or payment until all pass"),
        ("P1 - launch quality", "Custom domain; accessibility review; monitoring/alerts; onboarding; analytics consent; refund workflow; beta feedback", "Required for controlled public beta"),
        ("P2 - scale", "Enterprise SSO; external AI provider controls; service targets; advanced billing/tax; credential partnerships", "Complete as demand and risk grow"),
    ], [1700, 5200, 2460])

    doc.add_heading("Recommended release sequence", level=2)
    numbers(doc, [
        "Private production hardening: durable data, backups, secrets, monitoring, rate limits, legal drafts, and security testing.",
        "Invitation-only beta: 20-50 adults, free access, support-led onboarding, no public advertising, weekly issue review.",
        "Public free launch: open registration with clear legal acceptance, privacy controls, support, and capacity monitoring.",
        "Paid pilot: hosted checkout, self-service cancellation, receipts, refund process, tax review, and limited paid cohort.",
        "General paid launch: publish validated pricing, annual plans, team administration, service commitments, and mature operations.",
    ])

    new_chapter(doc, "1. Current product and launch scope", "Define exactly what the public is buying before writing legal copy, pricing, or marketing claims.")
    doc.add_heading("Public product description", level=2)
    para(doc, "SAP World is an independent enterprise simulation and tutoring platform for learning connected SAP S/4HANA business processes through guided transactions, troubleshooting, simulations, evidence coaching, capstones, and portfolio review.")
    app_ref(doc, "Live product surfaces: Sign-in -> Enterprise overview -> Learning centre -> Transaction tutor -> Portfolio")
    doc.add_heading("Claims you can make carefully", level=2)
    bullets(doc, [
        "Practice connected SAP process reasoning through a simulated enterprise.",
        "Follow guided pathways across eight end-to-end processes and multiple modules.",
        "Build a learning portfolio from saved progress, diagnostics, evidence notes, and capstone outcomes.",
        "Use deterministic simulations and document chains to connect operations with financial and management outcomes.",
    ])
    doc.add_heading("Claims to avoid without formal authority and evidence", level=2)
    bullets(doc, [
        "Official SAP training, official SAP certification, SAP-approved, SAP partner, or affiliated with SAP SE.",
        "Guaranteed employment, promotion, salary, certification pass, audit compliance, or production competence.",
        "Exact reproduction of every SAP S/4HANA screen, configuration, country localisation, or customer implementation.",
        "AI answers that are always correct, complete, or suitable for professional decisions.",
    ])
    warning(doc, "Use an independent-brand disclaimer in the site footer, Terms, About page, course pages, certificates/exports, and marketing material. Ask an IP solicitor to review the product name and all SAP mark usage before a paid launch.")

    new_chapter(doc, "2. P0 technical work before real users", "The public site must preserve accounts, progress, evidence, decisions, and audit records across deployments and server restarts.")
    doc.add_heading("2.1 Durable production data", level=2)
    numbers(doc, [
        "Provision managed PostgreSQL in a UK or appropriate regional location and obtain the production DATABASE_URL.",
        "Add DATABASE_URL only to Vercel Production environment variables; do not commit it to Git or expose it to the browser.",
        "Deploy the existing aggregate migration and verify all repositories use PostgreSQL rather than local JSON.",
        "Create test learner/admin records, progress, simulation, decision, mentor, and certification data.",
        "Redeploy and restart; confirm every record survives and /api/health reports durable storage.",
        "Enable automated backups and point-in-time recovery where offered. Document retention and perform a restore drill.",
        "Create a database access matrix: production application, named administrators, backup operator, and emergency access only.",
    ])
    app_ref(doc, "Public health check: /api/health; Admin evidence: Control plane -> Deployment gate summary")
    callout(doc, "Acceptance test", "After a full redeploy, the same learner can sign in and see identical progress, evidence, simulations, approval history, and certification state; the admin readiness gate no longer fails durable storage.", "EEF7F1", GREEN)

    doc.add_heading("2.2 Authentication and identity", level=2)
    bullets(doc, [
        "Keep passwords hashed using the existing server-side implementation; never log passwords or reset tokens.",
        "Set secure, HttpOnly, SameSite session cookies and rotate/revoke sessions after suspension, password reset, and privilege change.",
        "Rate-limit sign-in, registration, password recovery, mentor, simulation generation, and admin mutation endpoints.",
        "Add bot protection to registration and recovery if abuse appears; preserve an accessible alternative.",
        "Connect transactional recovery email with a verified sending domain, short token expiry, single use, and generic responses that do not reveal whether an email exists.",
        "Require multi-factor authentication for production administrators, Vercel, GitHub, database, email, payment, and domain accounts.",
        "Use OIDC/SAML for enterprise customers when sold; keep app roles and least-privilege authorisation server-side.",
    ])
    app_ref(doc, "Sign-in / Forgot password; Control plane -> Identity management / Access lifecycle")

    doc.add_heading("2.3 API and application safety", level=2)
    table(doc, ["Control", "Required implementation", "Evidence"], [
        ("Authorisation", "Check learner/admin role in every protected route; prevent object-ID access to another learner's records.", "Automated negative API tests"),
        ("Validation", "Validate type, length, format, allowed values, and ownership on every server input.", "Schema tests and rejected payload logs"),
        ("CSRF", "Protect state-changing cookie-authenticated requests with origin checks and an anti-CSRF design.", "Cross-site request test"),
        ("XSS", "Escape user content; do not render raw HTML from notes, capstones, mentor responses, or admin fields.", "Security test and code review"),
        ("Rate limits", "Limit costly, sensitive, and abuse-prone routes by account and network signal.", "429 tests and alert"),
        ("Secrets", "Keep database, email, AI, payment, and webhook secrets in environment variables; rotate on exposure.", "Secret inventory and rotation log"),
        ("Dependencies", "Run lockfile audit and update policy; remediate exploitable production vulnerabilities.", "CI report and owner"),
        ("Logs", "Exclude credentials, tokens, raw notes, capstone responses, payment data, and unnecessary personal data.", "Log sampling review"),
    ], [1600, 5200, 2560])

    doc.add_heading("2.4 Infrastructure, firewall, and observability", level=2)
    bullets(doc, [
        "Keep HTTPS enforced. Review the existing security headers and introduce a tested Content Security Policy that allows only required sources.",
        "Configure Vercel WAF/custom rules for admin endpoints, abusive clients, and high-risk traffic; start in log mode where blocking risk exists.",
        "Enable production runtime error, latency, traffic, function-cost, and firewall alerts with a named on-call owner.",
        "Send structured security and application events to a durable log service with retention and access control.",
        "Create service-level indicators: availability, error rate, sign-in success, database latency, mentor failures, and payment webhook failures.",
        "Document rollback: identify last known good deployment, database compatibility, restore steps, decision owner, and communication template.",
    ])
    app_ref(doc, "Control plane -> Observability / Deployment gate / Release operations; Vercel -> Observability / Firewall")

    new_chapter(doc, "3. Privacy and data protection", "Build the data map first; the privacy notice must describe the real system, not an aspirational template.")
    doc.add_heading("3.1 Create the data inventory", level=2)
    table(doc, ["Data", "Purpose", "Likely retention decision"], [
        ("Account identity and role", "Authentication, support, access control", "Account life plus a defined closure period"),
        ("Sessions and recovery", "Security and account recovery", "Short operational/security period"),
        ("Learning progress and evidence", "Deliver learning, portfolio, assessment", "Account life; provide deletion/export policy"),
        ("Simulations and ledgers", "Provide saved scenario learning", "Account life or user-controlled deletion"),
        ("Mentor prompts and responses", "Answer questions, quality and safety", "Minimise; short default unless user saves"),
        ("Admin/audit decisions", "Security, governance, certification review", "Longer justified period with restricted access"),
        ("Technical logs", "Security, debugging, reliability", "Short risk-based period; redact personal content"),
        ("Billing and invoices", "Contract, payment, tax, accounting", "Statutory/accounting period confirmed by adviser"),
        ("Marketing preferences", "Consent and communications", "Until withdrawal plus suppression evidence"),
    ], [2200, 3400, 3760])

    doc.add_heading("3.2 Privacy notice checklist", level=2)
    bullets(doc, [
        "Controller's legal name, trading name, postal address, privacy email, and representative/DPO details if applicable.",
        "Each category of personal data, source, purpose, and lawful basis; distinguish service delivery, security, analytics, marketing, AI, and legal obligations.",
        "Recipients/processors: Vercel, database, email, AI, analytics, support, payment, accounting, and professional advisers as actually used.",
        "International transfers, safeguards, and where information is processed.",
        "Retention criteria, deletion, backup expiry, and what must be retained by law or for claims/security.",
        "Rights and how to exercise access, correction, deletion, restriction, objection, portability, and consent withdrawal where applicable.",
        "Right to complain to the ICO, whether data is required for the contract, and consequences of not supplying it.",
        "Meaningful information about automated assessment or profiling if decisions significantly affect people; keep human review for certification decisions.",
        "Plain-language AI explanation: what context the mentor uses, whether prompts are sent to an external provider, and user safety instructions.",
    ])
    app_ref(doc, "Required public pages: footer -> Privacy; just-in-time notice at Register, Mentor, Capstone, Certification export, and Checkout")

    doc.add_heading("3.3 Cookies, analytics, and marketing", level=2)
    numbers(doc, [
        "Audit cookies, local storage, pixels, analytics scripts, and device-access technologies before adding a banner.",
        "Classify strictly necessary items separately from analytics, personalisation, advertising, and marketing.",
        "Do not load non-essential technologies before valid consent. Provide equally clear Accept and Reject choices and granular settings.",
        "Record consent, allow withdrawal at any time, and retest after adding integrations.",
        "Keep service messages separate from marketing. Obtain and record the correct permission for promotional email and include an easy unsubscribe.",
    ])
    warning(doc, "A cookie banner does not make tracking lawful by itself. The implementation must actually block non-essential technologies until the user chooses.")

    doc.add_heading("3.4 Governance actions", level=2)
    bullets(doc, [
        "Use the ICO fee self-assessment and pay/register if required.",
        "Maintain records of processing, processor contracts, subprocessor list, retention schedule, and rights-request procedure.",
        "Complete a data protection impact assessment before high-risk profiling, extensive monitoring, children's access, or sensitive AI use.",
        "Keep a breach log. Assess every personal-data incident; report a notifiable breach to the ICO without undue delay and, where feasible, within 72 hours.",
        "Decide the minimum age. For a simple first launch, restrict accounts and purchases to adults (18+) unless a specialist review supports younger users and parental/school arrangements.",
    ])

    new_chapter(doc, "4. Legal document pack", "Publish these documents through the footer and show the relevant version at registration and checkout.")
    table(doc, ["Document", "Must cover", "Launch point"], [
        ("Terms of Service", "Eligibility; account; licence; acceptable use; fees; cancellation; IP; disclaimers; suspension; liability; law; changes", "Before public registration"),
        ("Privacy Notice", "Real data map, purposes, lawful bases, processors, transfers, retention, rights, ICO", "Before collecting any public data"),
        ("Cookie Notice / settings", "Technologies, providers, purposes, duration, consent, withdrawal", "Before non-essential tracking"),
        ("Acceptable Use Policy", "No abuse, unlawful content, attacks, scraping, credential sharing, harassment, IP infringement", "Before open registration"),
        ("Subscription terms", "Price/tax, billing period, trial, renewal, reminders, cancellation, refunds, access after cancellation", "Before taking payment"),
        ("Refund and cancellation policy", "Cooling-off and digital service rules, refund method/timing, support route", "Before taking payment"),
        ("AI and education disclaimer", "Mentor limitations; verify answers; no professional advice; simulation; no guaranteed outcome", "Before mentor and in Terms"),
        ("Trademark / affiliation notice", "Independent product; ownership of SAP marks; no endorsement or official credential", "Footer, About, Terms, exports"),
        ("Accessibility statement", "Target standard, known issues, contact, response, improvements", "Public beta"),
        ("Subprocessor list", "Provider, purpose, location/transfer, change notice", "Enterprise readiness"),
        ("Business information", "Legal/trading name, geographic address, email, company/VAT number if applicable", "Footer, checkout, receipts"),
    ], [2100, 5100, 2160])

    doc.add_heading("Terms clause inventory", level=2)
    bullets(doc, [
        "Contract identity and acceptance: who operates SAP World, which version applies, and when a contract begins.",
        "Eligibility and account rules: adult launch, accurate details, unique account, security, no sharing, notice of compromise.",
        "Service licence: personal, limited, non-exclusive, non-transferable learning use; enterprise terms for teams.",
        "User content: learner retains ownership; grants only the licence needed to host, process, assess, and display it; prohibits confidential or unlawful uploads.",
        "Acceptable use and enforcement: investigation, proportionate suspension, appeal/contact, and urgent security action.",
        "Paid terms: complete price including tax, billing cycle, trial conversion, renewal, cancellation, refunds, failed payments, and access end date.",
        "Learning/AI disclaimers: simulation, no official SAP credential, no professional advice, verify outputs, no employment outcome.",
        "Intellectual property and feedback: ownership of platform/content, third-party marks, limited feedback licence, infringement contact.",
        "Availability and changes: maintenance, feature changes, beta status, data export/closure, and material change notice.",
        "Liability and consumer rights: preserve rights that cannot be excluded; solicitor drafts any cap/exclusion and governing-law wording.",
    ])
    callout(doc, "Legal review brief", "Give the solicitor the live app, this handbook, data inventory, processor list, planned countries, age limit, pricing, refund design, AI provider design, and every marketing claim. Ask for consumer, privacy, IP/trademark, and subscription review together.", LIGHT_BLUE)

    new_chapter(doc, "5. Consumer protection and subscription design", "Make joining, paying, renewing, and leaving equally understandable.")
    bullets(doc, [
        "Show the full price, tax treatment, billing period, minimum term, renewal, trial conversion, cancellation method, and service description before the customer pays.",
        "Use a clear final button such as 'Pay GBP X per month' rather than ambiguous wording.",
        "Send a durable confirmation email containing the contract, price, plan, date, cancellation/refund information, and support details.",
        "Provide a simple self-service cancellation route inside the account; do not force a phone call or hide it behind multiple screens.",
        "Send renewal/reminder communications required by applicable subscription rules and good practice; keep evidence of delivery.",
        "Avoid scarcity pressure, preselected add-ons, obstructive cancellation, disguised adverts, or interface choices that manipulate consent.",
        "Map the 14-day cancellation rules and any service/digital-content exception with a UK consumer solicitor before launch.",
        "Publish and operate a complaint process with response targets and escalation.",
    ])
    app_ref(doc, "Future paid flow: Pricing -> Plan summary -> Checkout -> Confirmation -> Account -> Billing / Cancel")

    new_chapter(doc, "6. Payments, tax, and financial controls", "Use a hosted payment surface so SAP World does not collect or store card numbers.")
    doc.add_heading("Recommended payment architecture", level=2)
    numbers(doc, [
        "Create products and recurring prices in Stripe using test mode first.",
        "Redirect the user to Stripe-hosted Checkout for payment.",
        "Receive signed webhooks server-side and process them idempotently; never grant paid access from the browser redirect alone.",
        "Store only provider/customer/subscription identifiers, entitlement, plan, status, and invoice references needed by the app.",
        "Offer the Stripe-hosted customer portal for invoices, payment method updates, plan changes, and cancellation.",
        "Reconcile webhook events, subscription state, refunds, chargebacks, and accounting records daily during launch.",
    ])
    warning(doc, "Never store card number, CVC, or full payment credentials in SAP World logs or database. Complete the payment provider's PCI compliance steps and maintain your own secure integration.")

    doc.add_heading("Tax and business setup", level=2)
    bullets(doc, [
        "Choose and register the operating entity; obtain a dedicated business bank account and bookkeeping process.",
        "Ask an accountant when UK VAT registration is required or beneficial and how online education/digital-service supplies are classified.",
        "If selling to consumers outside the UK, check place-of-supply, VAT/GST/sales-tax, invoice, consumer, privacy, and representative requirements before enabling that country.",
        "Display tax-inclusive consumer prices where required and issue compliant receipts/invoices.",
        "Maintain revenue, refund, fee, tax, payout, and deferred-revenue records; separate product analytics from accounting evidence.",
    ])

    new_chapter(doc, "7. Copyright, trademarks, and content provenance", "The product's learning value should come from original explanation and simulation, not copied SAP courseware or customer material.")
    doc.add_heading("Before public launch", level=2)
    bullets(doc, [
        "Run a UK trade-mark search for the product name and related classes; obtain legal advice on the risk of 'SAP World' before investing in marketing.",
        "Create a register for every text, image, icon, font, dataset, code library, and reference: owner, licence, source, attribution, and commercial-use status.",
        "Remove SAP logos, screenshots, manuals, course questions, or proprietary content unless a licence clearly permits the exact commercial use.",
        "Use SAP product names only as needed to describe compatibility or subject matter, following current SAP trademark guidance and legal advice.",
        "Add a notice such as: 'SAP and SAP S/4HANA are trademarks or registered trademarks of SAP SE or its affiliates. SAP World is an independent learning simulation and is not affiliated with, sponsored, endorsed, or certified by SAP SE.' Have counsel approve final wording.",
        "Require contributors to warrant originality/permissions and assign or license rights to the operating entity.",
        "Publish an IP complaint/takedown contact and keep a documented review and response process.",
        "Register your own distinctive brand mark if clearance supports it; Companies House and domain registration do not themselves create trade-mark rights.",
    ])
    app_ref(doc, "Review all learner pages, footer, About, Terms, exports, blueprints, tutor content, and marketing assets")

    new_chapter(doc, "8. Frontend safety, accessibility, and trust", "A safe public interface prevents mistakes, explains consequences, and works for people using keyboards, screen readers, zoom, and mobile devices.")
    table(doc, ["Area", "Public-launch standard"], [
        ("Accessibility", "Target WCAG 2.2 AA; keyboard flow, visible focus, names/labels, contrast, zoom/reflow, error identification, reduced motion, and screen-reader testing."),
        ("Forms", "Persistent labels, clear required fields, safe autocomplete, inline errors, preserved input, confirmation before destructive actions."),
        ("Consent", "No preselected optional tracking/marketing; Accept and Reject are equally clear; settings remain accessible."),
        ("Account safety", "Show active role and account menu; require reauthentication for sensitive changes; warn before data deletion."),
        ("User content", "Plain text by default; length limits; escaping; privacy reminders near mentor, evidence, capstone, and support inputs."),
        ("Errors", "Do not reveal stack traces, secrets, internal IDs, account existence, or another user's data."),
        ("Responsive design", "Test sign-in, all navigation, tutor steps, tables, dialogs, mentor, and admin actions on small mobile widths."),
        ("Trust", "Visible operator, support, policies, pricing, cancellation, security contact, and service status."),
    ], [2000, 7360])
    app_ref(doc, "Public pages plus every flow from Register through Portfolio; admin testing through Control plane")

    doc.add_heading("Accessibility verification", level=2)
    bullets(doc, [
        "Automated accessibility scan for every major view.",
        "Keyboard-only journey: registration, sign-in, navigation, tutor, forms, mentor, simulation, portfolio, and cancellation.",
        "Screen-reader checks with at least one Windows and one mobile combination.",
        "200% and 400% zoom/reflow; no hidden actions or horizontal scrolling for ordinary content.",
        "Colour contrast, focus visibility, status announcements, error summaries, and non-colour status cues.",
        "Publish known issues and a contact route; track remediation with owners and dates.",
    ])

    new_chapter(doc, "9. AI Mentor safety and quality", "The mentor must remain a learning aid with traceable context, minimal data, and honest limitations.")
    bullets(doc, [
        "Keep the current grounded local mentor as a safe baseline and label the provider/model when an external service is connected.",
        "Send only the minimum learner and simulation context needed; exclude raw credentials, unnecessary personal data, and private admin evidence.",
        "Contract with the AI provider for data use, retention, training, security, subprocessors, location, deletion, and incident notice.",
        "Filter or refuse requests involving credential theft, unauthorised access, harmful instructions, or sensitive personal/professional decisions.",
        "Persist conversations only with a stated purpose and retention period; provide deletion/export where appropriate.",
        "Run quality checks for grounding, citations/context, unsupported certainty, privacy, and pedagogical usefulness.",
        "Let users report a poor answer; keep a review queue and versioned evaluation set across all eight processes.",
        "Never allow an AI answer alone to approve certification, change an admin role, release content, or promote production.",
    ])
    app_ref(doc, "Ask SAP Mentor; Control plane -> Mentor readiness / Observability")

    new_chapter(doc, "10. Monetisation strategy", "Start with a narrow paid promise: guided practice plus credible evidence, not 'all of SAP'.")
    doc.add_heading("Recommended offer ladder", level=2)
    table(doc, ["Offer", "Who it serves", "Included value", "Pricing hypothesis"], [
        ("Free", "Curious learners", "Account, overview, limited pathways, sample mentor, starter portfolio", "GBP 0"),
        ("Individual Pro", "Career changers and consultants", "All pathways, diagnostics, simulations, advanced labs, capstones, exports", "Test GBP 15-25/month or GBP 150-240/year"),
        ("Cohort", "Learners needing structure", "Pro plus timetable, live clinics, feedback, and peer accountability", "Test GBP 199-499 per cohort"),
        ("Team", "Employers and training providers", "Seat management, assigned pathways, team analytics, review workflow", "Test GBP 25-40/user/month with minimum seats"),
        ("Enterprise", "Large organisations", "SSO, provisioning, private content, support, audit/security pack, negotiated terms", "Annual quote after discovery"),
        ("Assessment service", "Learners or organisations", "Human review of capstone evidence and feedback; not official SAP certification", "Separate per-review fee"),
    ], [1500, 2100, 3500, 2260])
    callout(doc, "Pricing rule", "These are test ranges, not final recommendations. Interview users, run willingness-to-pay tests, and measure conversion and retention before committing to annual public pricing.", LIGHT_BLUE)

    doc.add_heading("Monetisation opportunities", level=2)
    bullets(doc, [
        "Role-based learning bundles: procurement, finance, supply chain, manufacturing, quality, maintenance, and HR.",
        "Industry specialisation packs built from the ten blueprints and specialist labs.",
        "Employer onboarding and reskilling cohorts with assigned practice and evidence review.",
        "Training-provider licensing and white-labelled cohorts, subject to brand/IP terms.",
        "Human mentor or assessor add-ons with published standards and conflict controls.",
        "Enterprise content authoring, private scenarios, SSO, audit exports, and implementation support.",
        "Partnership referrals only when transparent, relevant, and clearly labelled; do not compromise learning advice.",
    ])

    doc.add_heading("What not to monetise early", level=2)
    bullets(doc, [
        "Advertising based on learner profiling or selling learner data.",
        "Pay-to-pass assessment, unverifiable credentials, or misleading SAP affiliation.",
        "Complex per-transaction credits that make learning unpredictable.",
        "Enterprise promises before backup, SSO, security review, support, and data-processing terms are operational.",
    ])

    doc.add_heading("Core metrics", level=2)
    table(doc, ["Funnel", "Metric", "Healthy question"], [
        ("Acquire", "Qualified visitor to registration", "Does the public description attract the intended learner?"),
        ("Activate", "First guided step and evidence note within 24 hours", "Does onboarding produce an early success?"),
        ("Learn", "Guided, diagnostic, evidence, and capstone completion", "Are users developing demonstrable capability?"),
        ("Retain", "Weekly active learners and pathway continuation", "Do users return for a meaningful next action?"),
        ("Convert", "Free-to-paid and trial-to-paid", "Is the paid promise clear and valuable?"),
        ("Sustain", "Revenue, refunds, chargebacks, gross margin, support cost", "Is growth operationally and financially healthy?"),
        ("Trust", "Incidents, complaints, accessibility issues, mentor reports", "Is the service becoming safer as it grows?"),
    ], [1400, 3100, 4860])

    new_chapter(doc, "11. Public launch operations", "A public launch is an operating capability, not simply changing a Vercel setting.")
    doc.add_heading("Domain, email, and public identity", level=2)
    bullets(doc, [
        "Clear the brand name before buying or promoting a custom domain.",
        "Connect the custom domain to Vercel, enforce HTTPS, and redirect alternate hostnames consistently.",
        "Create support@, privacy@, security@, and billing@ addresses with named owners and backup access.",
        "Configure SPF, DKIM, and DMARC for transactional and marketing mail; separate marketing where useful.",
        "Publish legal entity/business details, support hours, response targets, and a status page or incident banner process.",
    ])

    doc.add_heading("SEO and public content", level=2)
    bullets(doc, [
        "Add unique page titles, descriptions, canonical URL, robots rules, sitemap, social preview, favicon, and structured data where accurate.",
        "Keep authenticated learner/admin data out of search indexes and public metadata.",
        "Publish an honest landing page explaining audience, learning outcomes, product limitations, price, and sample workflow.",
        "Create support articles using the companion User Manual and link to them from the relevant app areas.",
        "Review every testimonial, outcome, comparison, and urgency claim for evidence and consumer-law fairness.",
    ])

    doc.add_heading("Support and incident response", level=2)
    numbers(doc, [
        "Create categories: account/access, progress/data, billing/refund, content accuracy, privacy rights, security, accessibility, and abuse.",
        "Set severity and response targets; privacy/security incidents bypass the ordinary queue.",
        "Train support not to request passwords, reset tokens, full card data, or unnecessary identity documents.",
        "Provide identity-verification steps for sensitive account requests.",
        "Maintain incident roles, communication templates, evidence preservation, containment, recovery, post-incident review, and regulator/customer notification decisions.",
    ])

    new_chapter(doc, "12. Testing and launch acceptance", "Launch only when a named owner signs the evidence for each critical journey.")
    table(doc, ["Test suite", "Minimum launch journeys"], [
        ("Functional", "Register; sign in/out; recover; suspend/reactivate; learn; evidence; diagnostic; capstone; simulation; approval; export; admin decision"),
        ("Persistence", "Create data, restart/redeploy, verify all records, concurrent writes, backup and restore"),
        ("Security", "Authorisation, IDOR, CSRF, XSS, rate limits, session invalidation, secret/log review, dependency scan"),
        ("Privacy", "Data map matches product, cookie blocking, consent withdrawal, export/delete/rights procedure, processor contracts"),
        ("Accessibility", "Automated plus keyboard, screen reader, zoom/reflow, contrast, mobile"),
        ("Compatibility", "Current Chrome, Edge, Firefox, Safari; desktop/tablet/mobile; slow network"),
        ("Performance", "Public landing, sign-in, dashboard, tutor, simulation, analytics, APIs under expected beta load"),
        ("Payments", "Test checkout, signed/idempotent webhooks, trial, renewal, failed payment, cancellation, refund, invoice, entitlement"),
        ("Recovery", "Rollback deployment, restore database, rotate secret, disable compromised account, communicate incident"),
    ], [2000, 7360])
    app_ref(doc, "Learner flows across the left navigation; admin evidence in Control plane; public /api/health")

    doc.add_heading("Final P0 sign-off checklist", level=2)
    bullets(doc, [
        "Production PostgreSQL is durable; migration, restart, backup, and restore tests passed.",
        "No critical/high exploitable security findings remain; admin MFA and least privilege are active.",
        "Terms, Privacy, Cookie, Acceptable Use, disclaimers, refund/subscription terms, and business details are published and legally reviewed.",
        "SAP/product-name trade-mark risk and all content licences have been reviewed.",
        "Recovery email, support, privacy, security, and billing contacts work.",
        "Cookie consent blocks non-essential tracking; marketing consent and unsubscribe work.",
        "Core learner and admin journeys pass mobile, console, accessibility, cross-browser, and performance tests.",
        "Monitoring, alerts, incident response, rollback, and on-call ownership are operational.",
        "If payments are enabled: hosted checkout, webhooks, cancellation, refunds, tax, receipts, and reconciliation pass end to end.",
        "The current Admin Control plane readiness fingerprint has an authorised release decision.",
    ])

    new_chapter(doc, "13. 90-day launch plan", "Keep the first public release deliberately small enough to observe and correct.")
    table(doc, ["Period", "Product and engineering", "Business, legal, and market"], [
        ("Days 1-14", "PostgreSQL, backups, restart test, email, rate limits, logs, alerts, critical security fixes", "Entity/name clearance, legal brief, data map, processor list, support addresses"),
        ("Days 15-30", "Legal UI, consent, deletion/export procedure, accessibility fixes, browser/performance test", "Solicitor review, accountant/tax review, beta terms, recruit 20-50 adult testers"),
        ("Days 31-45", "Invitation beta, issue triage, mentor evaluations, restore/rollback drill", "Weekly interviews, support metrics, refine positioning and onboarding"),
        ("Days 46-60", "Public free launch, capacity controls, analytics, documentation", "Publish landing/support content, measure activation/retention, no paid claims yet"),
        ("Days 61-75", "Stripe test/live setup, entitlement, portal, webhooks, refund/cancellation flows", "Final paid terms, price research, tax settings, paid pilot recruitment"),
        ("Days 76-90", "Paid pilot, reliability and billing review, top defects", "Evaluate conversion, retention, refunds, margin, trust; decide general paid launch"),
    ], [1200, 4100, 4060])

    doc.add_heading("Owners to name", level=2)
    bullets(doc, [
        "Product owner and release approver",
        "Engineering/security owner and incident lead",
        "Database/backup owner",
        "Privacy lead and rights-request owner",
        "Content/IP owner and SAP-mark review owner",
        "Support and accessibility owner",
        "Billing/refund and finance/tax owner",
        "AI Mentor quality and safety owner",
    ])

    new_chapter(doc, "14. Official sources and review notes", "Sources were checked on 30 June 2026. Requirements change; recheck before launch and before entering a new market.")
    source(doc, "ICO - What privacy information should we provide?", "https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/individual-rights/the-right-to-be-informed/what-privacy-information-should-we-provide/", "Required privacy information and transparency content.")
    source(doc, "ICO - Cookies and similar technologies", "https://ico.org.uk/for-organisations/direct-marketing-and-privacy-and-electronic-communications/guide-to-pecr/cookies-and-similar-technologies/", "Clear information, active consent, and strictly necessary exception.")
    source(doc, "ICO - Data protection fee", "https://ico.org.uk/for-organisations/data-protection-fee/", "Fee/registration assessment for organisations processing personal information.")
    source(doc, "ICO - Personal data breaches", "https://ico.org.uk/for-organisations/report-a-breach/personal-data-breach/personal-data-breaches-a-guide/", "Breach records, risk assessment, and 72-hour notification rule where applicable.")
    source(doc, "ICO - Encryption", "https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/security/a-guide-to-data-security/encryption/", "Risk-based technical and organisational security measures.")
    source(doc, "GOV.UK - Online and distance selling", "https://www.gov.uk/online-and-distance-selling-for-businesses", "Pre-contract information, confirmation, cancellation, and online selling duties.")
    source(doc, "GOV.UK - Online Choice Architecture", "https://www.gov.uk/government/collections/online-choice-architecture", "Clear, honest design and risks from dark patterns/subscription traps.")
    source(doc, "Digital Markets, Competition and Consumers Act 2024", "https://www.legislation.gov.uk/ukpga/2024/13/contents", "Consumer and subscription framework; obtain advice on current commencement and implementation.")
    source(doc, "GOV.UK - VAT rules for digital services", "https://www.gov.uk/guidance/the-vat-rules-if-you-supply-digital-services-to-private-consumers", "Place-of-supply and digital-service considerations, including automated learning examples.")
    source(doc, "GOV.UK - Using somebody else's intellectual property", "https://www.gov.uk/using-somebody-elses-intellectual-property", "Permission/licensing and infringement risk.")
    source(doc, "GOV.UK - Search for a trade mark", "https://www.gov.uk/search-for-trademark", "UK clearance search starting point.")
    source(doc, "SAP - Guidelines for Publications About SAP", "https://www.sap.com/dam/site/corporate/legal/SAP_external_book_projects_E.pdf", "SAP's current trademark/publication guidance; legal review is still required.")
    source(doc, "W3C - Web Content Accessibility Guidelines 2.2", "https://www.w3.org/TR/WCAG22/", "International web accessibility standard and testable success criteria.")
    source(doc, "Ofcom - Online Safety enforcement guidance", "https://www.ofcom.org.uk/siteassets/resources/documents/online-safety/information-for-industry/illegal-harms/online-safety-enforcement-guidance.pdf", "Scope includes certain user-to-user services; reassess before adding public sharing, forums, or community content.")
    source(doc, "Stripe - Hosted Checkout", "https://docs.stripe.com/payments/checkout", "Hosted payment page and subscriptions architecture.")
    source(doc, "Stripe - Customer portal", "https://docs.stripe.com/customer-management", "Self-service payment, invoice, subscription, and cancellation management.")
    source(doc, "PCI Security Standards Council - Small Merchant Guide", "https://www.pcisecuritystandards.org/wp-content/uploads/2022/05/Small_Merchant_Guide_to_Safe_Payments.pdf", "Secure payment responsibilities and reducing card-data exposure.")
    source(doc, "Vercel - Security overview", "https://vercel.com/docs/security", "HTTPS, shared responsibility, firewall, and platform security controls.")
    source(doc, "Vercel - Firewall", "https://vercel.com/docs/vercel-firewall", "DDoS mitigation, WAF, custom rules, managed rulesets, and observability.")
    source(doc, "Vercel - Environment variables", "https://vercel.com/docs/environment-variables", "Environment-scoped secrets and redeployment behaviour.")

    doc.add_heading("Professional review questions", level=2)
    table(doc, ["Adviser", "Questions to resolve before paid launch"], [
        ("Solicitor", "Brand/SAP marks; Terms; privacy/cookies; subscriptions/cancellation/refunds; consumer rights; liability; age; countries; AI; accessibility; online-safety scope"),
        ("Accountant/tax adviser", "Entity; bookkeeping; VAT registration; digital education classification; overseas B2C/B2B tax; invoices; revenue recognition"),
        ("Security professional", "Threat model; auth/session/API review; dependency/secret/log checks; penetration test; backup/restore; incident response"),
        ("Accessibility specialist", "WCAG 2.2 AA audit with assistive technology and prioritised remediation"),
    ], [2000, 7360])

    path = OUT / "SAP_World_Public_Launch_and_Monetisation_Handbook.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_user_manual())
    print(build_launch_handbook())
